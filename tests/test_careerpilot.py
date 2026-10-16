import tempfile
import unittest
import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from careerpilot import analyze_career_fit, extract_resume_text, get_llm_settings, match_job, retrieve_evidence, score_requirement_evidence
from storage import delete_analysis, get_analysis, init_db, list_analyses, save_analysis


class CareerToolsTests(unittest.TestCase):
    def setUp(self):
        self.local_provider = patch.dict("os.environ", {"CAREERPILOT_PROVIDER": "local"})
        self.local_provider.start()

    def tearDown(self):
        self.local_provider.stop()

    def test_empty_gemini_model_uses_current_default(self):
        with patch.dict("os.environ", {
            "CAREERPILOT_PROVIDER": "gemini",
            "GEMINI_API_KEY": "test-key-only",
            "GEMINI_MODEL": "",
        }, clear=False):
            settings = get_llm_settings()
        self.assertTrue(settings["enabled"])
        self.assertEqual(settings["provider"], "gemini")
        self.assertEqual(settings["model"], "gemini-3.5-flash-lite")

    def test_gemini_busy_error_is_explained_as_local_fallback(self):
        class GeminiBusyError(Exception):
            status_code = 503

        with patch("openai.OpenAI", side_effect=GeminiBusyError("temporary overload")):
            report = analyze_career_fit(
                "Excel and SQL projects",
                "Junior data analyst who uses Excel and SQL.",
                "Junior Data Analyst",
                {"provider": "gemini", "enabled": True, "api_key": "test-key-only", "model": "gemini-3.8-flash"},
            )
        self.assertEqual(report["provider"], "Local tools (Gemini unavailable)")
        self.assertIn("temporarily busy", report["provider_issue"])
        self.assertIn("local tools only", report["provider_issue"])

    def test_ai_requirement_score_requires_exact_resume_evidence(self):
        requirements = [
            {"requirement": "Python projects", "priority": "required", "evidence_quote": "Built a Python dashboard for class.", "evidence_strength": "strong", "reason": "A project demonstrates use."},
            {"requirement": "SQL experience", "priority": "required", "evidence_quote": "Used SQL in production.", "evidence_strength": "strong", "reason": "This statement is not in the resume."},
            {"requirement": "Power BI", "priority": "preferred", "evidence_quote": "", "evidence_strength": "none", "reason": "No evidence provided."},
        ]
        score, checked = score_requirement_evidence(requirements, "PROJECTS\nBuilt a Python dashboard for class.")
        self.assertEqual(score, 43)  # One of two required items has a verified quote; preferred items count less.
        self.assertEqual(checked[0]["evidence_label"], "Strong evidence")
        self.assertEqual(checked[1]["evidence_quote"], "")
        self.assertEqual(checked[1]["evidence_label"], "Not shown")

    def test_text_resume_is_extracted(self):
        self.assertIn("Python", extract_resume_text(b"Python developer", "resume.txt"))

    def test_docx_text_and_table_are_extracted(self):
        document = Document()
        document.add_paragraph("Analyst resume")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "SQL"
        table.cell(0, 1).text = "Excel"
        output = io.BytesIO()
        document.save(output)
        extracted = extract_resume_text(output.getvalue(), "resume.docx")
        self.assertIn("Analyst resume", extracted)
        self.assertIn("SQL", extracted)
        self.assertIn("Excel", extracted)

    def test_retrieval_finds_relevant_resume_passage(self):
        resume = "Built a Flask API in Python for a scheduling project.\nEnjoy hiking on weekends."
        evidence = retrieve_evidence(resume, "We need Python and Flask API experience.")
        self.assertIn("Flask API", evidence[0])

    def test_retrieval_removes_pdf_bullet_garbage_and_keeps_sections_separate(self):
        evidence = retrieve_evidence(
            "PROJECTS\nAI Chatbot\n\x7f\nBuilt a chatbot with Python.\n"
            "TECHNICAL SKILLS\n\x7f\nPython, Pandas, Excel",
            "Python analyst using Pandas and Excel",
        )
        self.assertTrue(evidence)
        self.assertTrue(all("\x7f" not in passage for passage in evidence))
        self.assertTrue(all(len(passage) <= 300 for passage in evidence))

    def test_tools_report_matches_and_gaps(self):
        report = analyze_career_fit(
            "Python developer with Flask and SQL project experience.",
            "Seeking Python, Flask, SQL, and AWS experience.",
        )
        self.assertEqual(set(report["matched_skills"]), {"python", "flask", "sql"})
        self.assertIn("aws", [gap["skill"] for gap in report["skill_gaps"]])
        self.assertTrue(report["interview_questions"])
        self.assertTrue(report["resume_review"])
        self.assertIn("python", report["tailored_resume"]["skills"])
        self.assertIn("aws", [gap["skill"] for gap in report["skill_gaps"]])

    def test_fit_score_uses_skill_priority_and_usage_evidence(self):
        job = "Required: Python. Required: SQL. Tableau is preferred."
        listed = match_job("TECHNICAL SKILLS\nPython, SQL, Tableau", job, "Analyst")
        evidenced = match_job("PROJECTS\nBuilt a Python report.\nUsed SQL to query records.", job, "Analyst")
        self.assertNotEqual(listed["match_score"], evidenced["match_score"])
        self.assertEqual(sum(item["weight"] for item in listed["score_breakdown"].values()), 100)
        self.assertIn("document comparison", listed["summary"])

    def test_fit_score_is_not_invented_when_job_has_no_detectable_requirements(self):
        report = match_job(
            "General background and community activities.",
            "We need someone to review customer records, write weekly summaries, and support the team.",
            "Office Assistant",
        )
        self.assertIsNone(report["match_score"])
        self.assertIn("could not estimate role fit", report["summary"])

    def test_configured_llm_score_uses_verified_requirements_not_model_percentage(self):
        payload = {
            "fit_assessment": {
                "rationale": "The project demonstrates Python. SQL is not supported by a resume example.",
                "requirements": [
                    {"requirement": "Python project experience", "priority": "required", "evidence_quote": "Built a Python report.", "evidence_strength": "strong", "reason": "A project example shows Python in use."},
                    {"requirement": "SQL experience", "priority": "required", "evidence_quote": "Used SQL in production.", "evidence_strength": "strong", "reason": "This invented quote is rejected."},
                    {"requirement": "Experience with APIs", "priority": "preferred", "evidence_quote": "", "evidence_strength": "none", "reason": "No resume evidence found."},
                ],
            },
            "summary": "A concise review.", "resume_review": [],
            "tailored_resume": {
                "headline": "Analyst", "summary": "Python project work.", "skills": ["python"],
                "highlights": ["Built a Python report."],
                "bullet_rewrites": [
                    {"source": "Built a Python report.", "suggestion": "Created a report with Python.", "reason": "Shows relevant project evidence."},
                    {"source": "Managed a 50 person team.", "suggestion": "Led a 50-person team.", "reason": "Should be rejected because it is not in the resume."},
                ],
                "note": "Review details.",
            },
        }
        fake_response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(payload)))])
        with patch.dict("os.environ", {"CAREERPILOT_PROVIDER": "openai", "OPENAI_API_KEY": "test-key"}), patch("openai.OpenAI") as client_type:
            client_type.return_value.chat.completions.create.return_value = fake_response
            report = analyze_career_fit("PROJECTS\nBuilt a Python report.", "Python and SQL analyst role", "Analyst")
        self.assertEqual(report["provider"], "OpenAI + RAG + local tools")
        self.assertEqual(report["match_score"], 43)  # One verified must-have example; unsupported criteria remain in the denominator.
        self.assertEqual(report["score_method"], "AI requirement review with verified resume quotations")
        self.assertEqual(len(report["requirements"]), 3)
        self.assertEqual(report["requirements"][1]["evidence_label"], "Not shown")
        schema = client_type.return_value.chat.completions.create.call_args.kwargs["response_format"]["json_schema"]["schema"]
        self.assertIn("requirements", schema["properties"]["fit_assessment"]["properties"])
        self.assertNotIn("score", schema["properties"]["fit_assessment"]["properties"])
        self.assertTrue(report["resume_review"])
        self.assertIn("python", report["tailored_resume"]["skills"])
        self.assertEqual(len(report["tailored_resume"]["bullet_rewrites"]), 1)
        self.assertEqual(report["tailored_resume"]["bullet_rewrites"][0]["source"], "Built a Python report.")

    def test_analysis_works_without_api_key(self):
        with patch.dict("os.environ", {"CAREERPILOT_PROVIDER": "local"}):
            report = analyze_career_fit("Python developer", "Python and AWS role")
        self.assertIn(report["provider"], {"Local tools", "LangGraph + local tools"})
        self.assertIn("aws", [gap["skill"] for gap in report["skill_gaps"]])

    def test_analysis_falls_back_when_langgraph_is_unavailable(self):
        with patch.dict("os.environ", {"CAREERPILOT_PROVIDER": "local"}), patch("careerpilot._langgraph_tools", side_effect=ImportError):
            report = analyze_career_fit("Python developer", "Python and AWS role")
        self.assertEqual(report["provider"], "Local tools")


class StorageTests(unittest.TestCase):
    def test_saved_analysis_appears_in_history_and_detail(self):
        with tempfile.TemporaryDirectory() as folder:
            db_path = Path(folder) / "test.db"
            init_db(db_path)
            report = {"match_score": 50, "summary": "A useful first pass."}
            analysis_id = save_analysis(
                db_path,
                job_title="Data Analyst",
                resume_name="resume.txt",
                resume_text="resume text",
                job_text="job text",
                report=report,
                created_at="2026-01-01T00:00:00+00:00",
            )
            self.assertEqual(get_analysis(db_path, analysis_id)["report"], report)
            self.assertEqual(list_analyses(db_path)[0]["job_title"], "Data Analyst")
            self.assertTrue(delete_analysis(db_path, analysis_id))
            self.assertIsNone(get_analysis(db_path, analysis_id))
            self.assertFalse(delete_analysis(db_path, analysis_id))


if __name__ == "__main__":
    unittest.main()
